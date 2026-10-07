import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def parse_markdown_to_docx(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        content = f.read()

    doc = Document()

    # Set Margins (2.54 cm / 1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Set base styles
    normal_style = doc.styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal_style.paragraph_format.line_spacing = 1.35
    normal_style.paragraph_format.space_after = Pt(6)

    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_block_lines = []
    code_lang = ""

    def add_formatted_runs(p, text):
        # Parses bold, italic, code backticks, etc.
        tokens = re.split(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|`.*?`|\$\$.*?\$\$|\$.*?\$)', text)
        for token in tokens:
            if not token:
                continue
            if token.startswith('***') and token.endswith('***'):
                run = p.add_run(token[3:-3])
                run.bold = True
                run.italic = True
            elif token.startswith('**') and token.endswith('**'):
                run = p.add_run(token[2:-2])
                run.bold = True
            elif token.startswith('*') and token.endswith('*'):
                run = p.add_run(token[1:-1])
                run.italic = True
            elif token.startswith('`') and token.endswith('`'):
                run = p.add_run(token[1:-1])
                run.font.name = 'Consolas'
                run.font.size = Pt(10.5)
                run.font.color.rgb = RGBColor(0x8B, 0x00, 0x00)
            elif (token.startswith('$$') and token.endswith('$$')) or (token.startswith('$') and token.endswith('$')):
                raw = token.strip('$')
                run = p.add_run(raw)
                run.italic = True
                run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            else:
                p.add_run(token)

    while i < len(lines):
        line = lines[i]

        # Code block handler
        if line.strip().startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lang = line.strip()[3:].strip()
                code_block_lines = []
            else:
                in_code_block = False
                # Add code block box / callout
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.right_indent = Inches(0.4)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.line_spacing = 1.15
                
                label = f"[{code_lang.upper()} DIAGRAM / SPESIFIKASI]" if code_lang else "[BLOK KODE]"
                run_lbl = p.add_run(label + "\n")
                run_lbl.bold = True
                run_lbl.font.size = Pt(9.5)
                run_lbl.font.color.rgb = RGBColor(0x55, 0x55, 0x77)

                code_text = "\n".join(code_block_lines)
                run_code = p.add_run(code_text)
                run_code.font.name = 'Consolas'
                run_code.font.size = Pt(9.5)
                run_code.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            i += 1
            continue

        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue

        stripped = line.strip()

        # Blank line
        if not stripped:
            i += 1
            continue

        # Markdown Table Detection
        if stripped.startswith('|') and stripped.endswith('|'):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].strip().endswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            
            if len(table_lines) >= 2:
                # Parse header and rows
                raw_rows = []
                for tl in table_lines:
                    # Ignore divider line (e.g. |---|---|)
                    if re.match(r'^\|[\s:-|-]+\|$', tl):
                        continue
                    cols = [c.strip() for c in tl.split('|')[1:-1]]
                    raw_rows.append(cols)
                
                if raw_rows:
                    num_cols = max(len(r) for r in raw_rows)
                    tbl = doc.add_table(rows=len(raw_rows), cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    set_table_borders(tbl)

                    for r_idx, row_data in enumerate(raw_rows):
                        for c_idx, cell_value in enumerate(row_data):
                            if c_idx < num_cols:
                                cell = tbl.cell(r_idx, c_idx)
                                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                                set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                                p = cell.paragraphs[0]
                                p.paragraph_format.line_spacing = 1.15
                                p.paragraph_format.space_after = Pt(2)
                                if r_idx == 0:
                                    set_cell_background(cell, "F0F4F8")
                                    run = p.add_run(cell_value)
                                    run.bold = True
                                    run.font.name = 'Times New Roman'
                                    run.font.size = Pt(10.5)
                                    run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
                                else:
                                    if r_idx % 2 == 1:
                                        set_cell_background(cell, "FAFAFA")
                                    add_formatted_runs(p, cell_value)
                                    for r in p.runs:
                                        r.font.name = 'Times New Roman'
                                        r.font.size = Pt(10)
                    doc.add_paragraph() # Spacing after table
            continue

        # Horizontal rule
        if stripped == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            # Add subtle divider line
            run = p.add_run('—' * 55)
            run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # Title (H1 starting with # )
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(14)
            run = p.add_run(stripped[2:].strip())
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(0x11, 0x22, 0x44)
            i += 1
            continue

        # Heading 2 (## )
        if stripped.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[3:].strip())
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            i += 1
            continue

        # Heading 3 (### )
        if stripped.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(stripped[4:].strip())
            run.bold = True
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12.5)
            run.font.color.rgb = RGBColor(0x22, 0x33, 0x55)
            i += 1
            continue

        # Bullet lists (- or *)
        if stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(4)
            add_formatted_runs(p, stripped[2:].strip())
            i += 1
            continue

        # Numbered lists (1. , 2. )
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.line_spacing = 1.25
            p.paragraph_format.space_after = Pt(4)
            add_formatted_runs(p, num_match.group(2).strip())
            i += 1
            continue

        # Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.35
        p.paragraph_format.space_after = Pt(6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_formatted_runs(p, stripped)
        i += 1

    doc.save(docx_path)
    print(f"File successfully created at: {docx_path}")

if __name__ == '__main__':
    md_file = r'c:\Users\Acer\Documents\AntiGravity\FundamenSainsData\Fundamen-Sains-Data\docs\artikel_domain_aplikasi_skripsiflow.md'
    docx_file = r'c:\Users\Acer\Documents\AntiGravity\FundamenSainsData\Fundamen-Sains-Data\docs\artikel_domain_aplikasi_skripsiflow.docx'
    parse_markdown_to_docx(md_file, docx_file)
